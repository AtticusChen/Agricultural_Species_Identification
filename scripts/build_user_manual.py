# ruff: noqa: E501

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "薇甘菊识别系统-零基础操作手册.docx"
SCREENSHOT = ROOT / "artifacts" / "ui" / "mikania-detection-page.png"

FONT = "Arial Unicode MS"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17372D"
MUTED = "66736E"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "EDF4E8"
LIGHT_GOLD = "FFF6D8"
LIGHT_RED = "FDECEC"
BORDER = "C8D2CD"
WHITE = "FFFFFF"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = TABLE_INDENT_DXA) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: float | None = None, *, bold=False, color: str | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 9, color=MUTED)


def add_numbering(doc: Document) -> tuple[list[int], int]:
    numbering = doc.part.numbering_part.element

    def next_id(tag: str, attr: str) -> int:
        values = [int(node.get(qn(attr))) for node in numbering.findall(qn(tag))]
        return max(values, default=0) + 1

    def define(marker: str, num_format: str, start: int) -> int:
        abstract_id = next_id("w:abstractNum", "w:abstractNumId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start_node = OxmlElement("w:start")
        start_node.set(qn("w:val"), str(start))
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_format)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), marker)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.extend([tabs, ind, spacing])
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        fonts.set(qn("w:eastAsia"), FONT)
        rpr.append(fonts)
        level.extend([start_node, fmt, lvl_text, justification, ppr, rpr])
        abstract.append(level)
        numbering.append(abstract)

        num_id = next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        numbering.append(num)
        return num_id

    decimal_ids = [define("%1.", "decimal", 1) for _ in range(6)]
    return decimal_ids, define("●", "bullet", 1)


def apply_num(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    ppr.append(num_pr)


def add_numbered(doc: Document, num_id: int, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    apply_num(paragraph, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, 11, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, 11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11)


def add_bullet(doc: Document, num_id: int, text: str, *, compact: bool = False) -> None:
    paragraph = doc.add_paragraph()
    apply_num(paragraph, num_id)
    if compact:
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.1
    set_run_font(paragraph.add_run(text), 10 if compact else 11)


def add_callout(doc: Document, label: str, text: str, fill: str = LIGHT_BLUE) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    ppr.append(borders)
    set_run_font(paragraph.add_run(f"{label}  "), 10.5, bold=True, color=DARK_BLUE)
    set_run_font(paragraph.add_run(text), 10.5, color=INK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(header), 10, bold=True, color=INK)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(value), 9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def setup_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("薇甘菊智能识别系统｜用户操作手册"), 8.5, color=MUTED)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("第 "), 9, color=MUTED)
    add_field(paragraph, "PAGE")
    set_run_font(paragraph.add_run(" 页"), 9, color=MUTED)


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("MIKANIA MICRANTHA · FIELD DETECTOR"), 10, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("薇甘菊智能识别系统"), 28, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    set_run_font(subtitle.add_run("零基础安装与使用操作手册"), 17, bold=True, color=BLUE)
    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.6)
    lead.paragraph_format.right_indent = Inches(0.6)
    set_run_font(
        lead.add_run("无需预装 Python｜解压后双击安装｜模型和图片均在本机处理"),
        11,
        color=MUTED,
    )
    for _ in range(6):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("便携版 v2.1.0 · 2026年8月"), 10.5, bold=True, color=INK)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(note.add_run("适用于 Windows、macOS 和 64 位 Linux 笔记本"), 9.5, color=MUTED)
    doc.add_page_break()


def build_manual() -> Path:
    doc = Document()
    setup_styles(doc)
    setup_header_footer(doc)
    decimal_ids, bullet_id = add_numbering(doc)
    windows_steps, macos_steps, linux_steps, usage_steps, sharing_steps, migration_steps = decimal_ids
    add_cover(doc)

    doc.add_heading("先看这里：只需要记住两次双击", level=1)
    add_callout(
        doc,
        "最简流程",
        "第一次使用：全部解压 → 双击“01-一键安装” → 双击“02-启动识别系统”。以后使用只需双击“02-启动识别系统”。",
        LIGHT_GREEN,
    )
    add_table(
        doc,
        ["步骤", "您要做什么", "看到什么算成功"],
        [
            ["1", "把 ZIP 压缩包全部解压到一个普通文件夹", "能看到以 00、01、02、03 开头的文件；Windows 还有 04 诊断文件"],
            ["2", "保持联网，双击与电脑系统对应的 01 文件", "窗口最后显示“安装成功”"],
            ["3", "双击与电脑系统对应的 02 文件", "浏览器自动出现识别页面"],
            ["4", "选择照片并点击“开始识别”", "页面显示判断结果和标注图"],
        ],
        [900, 4460, 4000],
    )
    add_callout(
        doc,
        "重要",
        "不能直接在 ZIP 压缩包预览窗口中运行。必须先选择“全部解压”或双击 ZIP 后把整个文件夹拖到桌面。",
        LIGHT_GOLD,
    )

    doc.add_heading("1. 安装前准备", level=1)
    doc.add_heading("1.1 电脑和网络要求", level=2)
    add_table(
        doc,
        ["项目", "最低要求", "说明"],
        [
            ["操作系统", "Windows 10/11；macOS 14+；Ubuntu 22.04+", "均要求 64 位系统"],
            ["内存", "4 GB 可用内存", "8 GB 或以上使用更流畅"],
            ["磁盘", "2 GB 可用空间", "运行环境安装在交付文件夹内部"],
            ["网络", "首次安装需要联网", "安装完成后可离线识别图片"],
            ["权限", "普通用户权限", "通常不需要管理员账户"],
        ],
        [1700, 3160, 4500],
    )
    doc.add_heading("1.2 解压后应该看到的文件", level=2)
    add_bullet(doc, bullet_id, "00-请先阅读-快速开始.txt：一页式快速说明。")
    add_bullet(doc, bullet_id, "01-一键安装-电脑系统：第一次使用时运行一次。")
    add_bullet(doc, bullet_id, "02-启动识别系统-电脑系统：日常使用入口。")
    add_bullet(doc, bullet_id, "03-启动局域网共享-电脑系统：让同一 Wi-Fi 下的其他设备访问。")
    add_bullet(doc, bullet_id, "04-Windows故障诊断.bat：Windows 安装或启动异常时运行，窗口不会闪退。")
    add_bullet(doc, bullet_id, "薇甘菊识别系统-零基础操作手册：本手册的 Word 和 PDF 版本。")
    add_bullet(doc, bullet_id, "程序源码：网页、推理、数据处理、训练代码和研发脚本。")
    add_bullet(doc, bullet_id, "模型文件：model.onnx 网页识别模型和 model.pt 训练权重。")
    add_bullet(doc, bullet_id, "runtime（安装与运行脚本）：环境安装器、启动器和安装日志，请不要移动或改名。")
    add_bullet(doc, bullet_id, "示例图片：可用于第一次测试系统。")

    doc.add_heading("2. Windows 实操步骤", level=1)
    add_numbered(doc, windows_steps, "右键 ZIP 压缩包，选择“全部解压”，再点击“提取”。", bold_prefix="右键 ZIP 压缩包")
    add_numbered(doc, windows_steps, "进入解压后的文件夹，双击“01-一键安装-Windows.bat”。")
    add_numbered(doc, windows_steps, "如果 Windows 显示安全提示，确认文件来自本项目交付包后，选择“更多信息”→“仍要运行”。")
    add_numbered(doc, windows_steps, "保持黑色安装窗口开启。程序会自动下载独立 Python 和识别组件，无需输入命令。")
    add_numbered(doc, windows_steps, "窗口显示“安装成功”后按任意键关闭。")
    add_numbered(doc, windows_steps, "双击“02-启动识别系统-Windows.bat”，等待浏览器自动打开。")
    add_callout(doc, "安装时间", "网络正常时通常需要 5 至 15 分钟。下载过程中窗口文字会滚动，这是正常现象。", LIGHT_BLUE)
    add_callout(doc, "出现异常", "新版窗口会停留显示错误。请双击“04-Windows故障诊断.bat”，拍下窗口内容，并保留 runtime/logs 中的安装日志。", LIGHT_GOLD)
    add_callout(doc, "不要做", "不要关闭正在安装的窗口；不要移动 runtime 文件夹；不要把交付包放在只读 U 盘中直接运行。", LIGHT_RED)

    doc.add_heading("3. macOS 实操步骤", level=1)
    add_numbered(doc, macos_steps, "双击 ZIP 完成解压，进入解压后的文件夹。")
    add_numbered(doc, macos_steps, "第一次安装时，右键“01-一键安装-macOS.command”，选择“打开”。")
    add_numbered(doc, macos_steps, "系统再次询问时点击“打开”。不要直接双击第一次运行，右键打开更容易通过安全确认。")
    add_numbered(doc, macos_steps, "保持终端窗口开启，等待显示“安装成功”。")
    add_numbered(doc, macos_steps, "以后双击“02-启动识别系统-macOS.command”即可。")
    add_callout(
        doc,
        "如果仍被拦截",
        "打开“系统设置”→“隐私与安全性”，在安全提示旁点击“仍要打开”，然后重新运行安装文件。",
        LIGHT_GOLD,
    )

    doc.add_heading("4. Linux 实操步骤", level=1)
    add_numbered(doc, linux_steps, "解压 ZIP，进入解压后的文件夹。")
    add_numbered(doc, linux_steps, "在文件管理器中右键 01 文件，打开属性并允许“作为程序执行”。")
    add_numbered(doc, linux_steps, "运行“01-一键安装-Linux.sh”，等待安装完成。")
    add_numbered(doc, linux_steps, "运行“02-启动识别系统-Linux.sh”。")
    add_callout(
        doc,
        "文件管理器不支持双击脚本时",
        "请让单位技术人员在该文件夹运行 01 和 02 文件。普通使用者无需理解脚本内容。",
        LIGHT_BLUE,
    )

    doc.add_heading("5. 日常识别操作", level=1)
    add_numbered(doc, usage_steps, "双击“02-启动识别系统”文件，并保持启动窗口不要关闭。")
    add_numbered(doc, usage_steps, "浏览器打开后，确认右上角显示“识别模型已就绪”。")
    add_numbered(doc, usage_steps, "点击照片区域，选择 JPG、PNG 或 WebP 植物照片，也可以把照片拖入页面。")
    add_numbered(doc, usage_steps, "通常保持“精细扫描（推荐）”；只在电脑较慢或需要快速初筛时选择“快速扫描”。")
    add_numbered(doc, usage_steps, "需要时调整“显示置信度”。精细扫描默认 15%，不确定时保持默认值。")
    add_numbered(doc, usage_steps, "点击“开始识别”，等待结果出现。")
    add_numbered(doc, usage_steps, "完成后可以点击“重新选择”继续识别下一张照片。")
    add_numbered(doc, usage_steps, "使用结束时关闭启动窗口，系统即停止运行。")
    if SCREENSHOT.is_file():
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(SCREENSHOT), width=Inches(5.55))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(caption.add_run("图1  识别页面、照片上传区域与检测结果示例"), 9, color=MUTED)

    doc.add_page_break()

    doc.add_heading("6. 怎样理解识别结果", level=1)
    add_table(
        doc,
        ["页面结果", "含义", "建议操作"],
        [
            ["识别到薇甘菊", "模型在照片中发现一个或多个疑似区域", "查看绿色框位置，必要时请专业人员复核"],
            ["未识别到薇甘菊", "当前阈值下没有检测到目标", "不能等同于确认不存在，可换清晰近景或降低阈值复查"],
            ["精细扫描", "把高清照片分区检查后合并结果", "适合现场照片和多区域覆盖场景，速度较慢但漏检更少"],
            ["快速扫描", "整张照片一次缩放后检查", "适合低配置电脑快速初筛，可能漏掉远处或较小区域"],
            ["最高置信度", "模型对最可信检测区域的相对把握", "用于排序和辅助判断，不是物种鉴定结论"],
            ["处理时间", "本机完成一次识别所用时间", "首张图较慢属于正常现象"],
        ],
        [2250, 3550, 3560],
    )
    add_callout(
        doc,
        "使用边界",
        "本系统用于现场初步筛查。照片过暗、过远、遮挡严重或与训练数据差异较大时可能漏检或误检，最终判断应结合专业人员意见。",
        LIGHT_GOLD,
    )

    doc.add_heading("7. 让同一 Wi-Fi 下的其他设备访问", level=1)
    add_numbered(doc, sharing_steps, "在安装系统的主电脑上运行对应的“03-启动局域网共享”文件。")
    add_numbered(doc, sharing_steps, "如果防火墙询问是否允许访问，请允许“专用网络”访问。")
    add_numbered(doc, sharing_steps, "查看启动窗口中显示的地址，例如 http://192.168.1.20:18080。")
    add_numbered(doc, sharing_steps, "让手机或另一台电脑连接同一个 Wi-Fi，在浏览器输入该地址。")
    add_callout(doc, "安全提醒", "局域网共享模式没有账号密码。只用于可信的单位或家庭局域网，不要直接暴露到互联网。", LIGHT_RED)

    doc.add_heading("8. 常见问题与解决办法", level=1)
    add_table(
        doc,
        ["现象", "最可能原因", "处理步骤"],
        [
            ["Windows 双击后闪退", "没有全部解压，或仍在使用旧版交付包", "使用 v2.1.0；对 ZIP 选择“全部解压”；运行 04-Windows故障诊断.bat"],
            ["网页显示无法访问", "启动窗口已关闭或系统尚未启动", "重新双击 02 文件；保持窗口开启；手动打开 http://127.0.0.1:18080"],
            ["安装一直没有完成", "网络慢或无法访问下载站点", "等待一段时间；换稳定网络；关闭后重新运行 01 文件"],
            ["提示端口被占用", "已经启动过一次或其他程序使用 18080", "关闭旧启动窗口和浏览器，再重新运行 02 文件"],
            ["提示模型校验失败", "压缩包未完整解压或模型损坏", "删除当前解压目录，从原始 ZIP 重新全部解压"],
            ["照片无法上传", "格式不支持或文件损坏", "换用 JPG、PNG、WebP；先用系统照片程序确认能打开"],
            ["识别很慢", "电脑性能较低或首次加载模型", "等待首张完成；关闭其他大型程序；后续图片通常更快"],
            ["漏检或检测不完整", "使用了快速扫描、目标较远或遮挡严重", "切换到精细扫描；使用清晰近景；保持默认 15% 后复查"],
            ["误把其他植物识别为薇甘菊", "相似叶形或复杂背景造成误报", "适当提高置信度；换角度补拍花、藤蔓和叶片；请专业人员复核"],
        ],
        [2350, 2840, 4170],
    )
    add_callout(
        doc,
        "联系技术支持时",
        "请同时提供：电脑系统版本、出现问题的步骤、04 诊断窗口照片，以及 runtime/logs 文件夹中的安装日志。这样可以最快定位问题。",
        LIGHT_BLUE,
    )

    doc.add_heading("9. 更新、迁移与卸载", level=1)
    doc.add_heading("9.1 换到另一台电脑", level=2)
    add_numbered(doc, migration_steps, "复制原始 v2.1.0 ZIP 到新电脑。不要复制已经安装完成的 runtime 文件夹。")
    add_numbered(doc, migration_steps, "在新电脑重新全部解压，并运行对应系统的 01 一键安装。")
    add_numbered(doc, migration_steps, "安装完成后运行 02 启动文件。")
    doc.add_heading("9.2 卸载", level=2)
    add_bullet(doc, bullet_id, "先关闭识别系统启动窗口。", compact=True)
    add_bullet(doc, bullet_id, "删除整个 mikania-portable-2.1.0 文件夹即可。", compact=True)
    add_bullet(doc, bullet_id, "系统不会卸载或修改电脑原有 Python。", compact=True)
    doc.add_heading("9.3 更新模型或软件", level=2)
    add_bullet(doc, bullet_id, "不要直接覆盖“程序源码”“模型文件”或 runtime 中的单个文件。", compact=True)
    add_bullet(doc, bullet_id, "使用新的完整交付包并重新安装，以避免版本混用。", compact=True)

    doc.add_heading("10. 安装与使用验收清单", level=1)
    for item in (
        "已经从 ZIP 全部解压，而不是在压缩包中直接运行。",
        "01 安装窗口最后显示“安装成功”。",
        "02 启动后浏览器自动打开识别页面。",
        "页面右上角显示“识别模型已就绪”。",
        "示例图片可以上传，并能得到识别结果。",
        "关闭启动窗口后网页停止，重新运行 02 后恢复。",
        "需要共享时，另一台同 Wi-Fi 设备能够打开窗口显示的局域网地址。",
    ):
        add_bullet(doc, bullet_id, f"□ {item}", compact=True)

    doc.add_heading("附录：交付说明", level=1)
    add_bullet(doc, bullet_id, "算法模型：YOLO26s 薇甘菊单类别目标检测模型。", compact=True)
    add_bullet(doc, bullet_id, "默认识别：精细扫描，置信度阈值 15%；可切换快速扫描。", compact=True)
    add_bullet(doc, bullet_id, "模型权重：model.onnx 用于 ONNX Runtime CPU 本机离线推理；model.pt 用于继续训练、评估和重新导出。", compact=True)
    add_bullet(doc, bullet_id, "默认本机地址：http://127.0.0.1:18080。", compact=True)
    add_bullet(doc, bullet_id, "支持图片：JPG、PNG、WebP。", compact=True)
    add_bullet(doc, bullet_id, "数据处理：上传照片仅在运行该系统的本机内处理。", compact=True)
    add_bullet(doc, bullet_id, "环境安装：使用项目目录内的独立 uv 和受管理 Python 3.11，不依赖系统 Python。", compact=True)
    add_bullet(
        doc,
        bullet_id,
        "许可提示：模型由 Ultralytics YOLO 训练和导出。若用于闭源商业产品或对外收费服务，应在发布前核查并取得适用的软件许可。",
        compact=True,
    )

    doc.core_properties.title = "薇甘菊识别系统零基础操作手册"
    doc.core_properties.subject = "便携版一键安装、启动、识别与故障处理"
    doc.core_properties.author = "薇甘菊智能识别项目"
    doc.core_properties.keywords = "薇甘菊,目标检测,用户手册,一键安装"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_manual())
